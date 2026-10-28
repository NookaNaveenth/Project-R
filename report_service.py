"""Generate attendance XLSX reports using the template structure from attendance_october_2026_updated.xlsx."""

import calendar
import os
import re
from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from attendance import get_connection, get_employee, get_employee_attendance_history

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Exact 9 fixed metadata columns matching the template (without erroneous "Working From")
FIXED_COLUMNS = [
    "#", "Emp-ID", "Last Name", "First Name", "Dept", "Practice",
    "Reporting Person", "Work Location", "HR Contact",
]
# Summary columns starting at column 41 (AO)
SUMMARY_COLUMNS = [
    "P", "Half Day", "Actual Working Days", "Tot. Leaves",
    "Holidays", "Week ends", "No of Days in a Month",
]


def _get_template_path():
    """Discover the Excel template path without hardcoding a single fixed location."""
    candidates = [
        r"C:\Users\mcconf\Downloads\attendance_october_2026_updated.xlsx",
        os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "attendance_template.xlsx"),
        os.path.join(os.path.dirname(os.path.abspath(__file__)), "attendance_template.xlsx"),
    ]
    for p in candidates:
        if os.path.exists(p):
            return p
    return None


def _load_or_create_workbook():
    """Load the template workbook if available, otherwise initialize a fresh Workbook."""
    template_path = _get_template_path()
    if template_path:
        try:
            wb = load_workbook(template_path)
            return wb, True
        except Exception:
            pass
    return Workbook(), False


def _parse_report_date(value):
    if not value:
        return date.today()
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()
    except (TypeError, ValueError) as error:
        raise ValueError("date must use YYYY-MM-DD format") from error


def _period_dates(period, selected_date=None, start_date=None, end_date=None):
    norm_period = str(period or "month").strip().lower()
    if norm_period in ("day", "daily"):
        target = _parse_report_date(selected_date or start_date)
        return target, target
    if norm_period in ("week", "weekly"):
        selected = _parse_report_date(selected_date or start_date)
        start = selected - timedelta(days=selected.weekday())
        return start, start + timedelta(days=6)
    if norm_period in ("month", "monthly"):
        selected = _parse_report_date(selected_date or start_date)
        start = selected.replace(day=1)
        end = selected.replace(day=calendar.monthrange(selected.year, selected.month)[1])
        return start, end
    if norm_period in ("custom", "custom_range", "range"):
        if not start_date or not end_date:
            raise ValueError("Both start_date and end_date are required for a custom date range.")
        start, end = _parse_report_date(start_date), _parse_report_date(end_date)
        if start > end:
            raise ValueError("Start date must be before or equal to end date.")
        return start, end
    raise ValueError("period must be 'day', 'week', 'month', or 'custom'")


def _history_for_range(employee_id, start_date, end_date):
    rows = {}
    cursor = start_date.replace(day=1)
    while cursor <= end_date:
        history = get_employee_attendance_history(employee_id, cursor.year, cursor.month)
        if history is None:
            raise LookupError(f"Employee '{employee_id}' not found")
        rows.update({item["date"]: item for item in history.get("attendance", [])})
        if cursor.month == 12:
            cursor = cursor.replace(year=cursor.year + 1, month=1)
        else:
            cursor = cursor.replace(month=cursor.month + 1)
    return [
        rows.get((start_date + timedelta(days=offset)).isoformat())
        for offset in range((end_date - start_date).days + 1)
    ]


def _employee_rows(department=None, employee_ids=None, practice=None):
    connection = get_connection()
    try:
        query = """SELECT employee_id, employee_name, shift_start, shift_end,
                          email, login_id, department, designation, team_name,
                          reports_to, reports_to_mail_id, hr_contact, hr_email,
                          work_location, is_flexible_shift, reporting_hr, practice, project
                   FROM employees WHERE 1=1"""
        params = []
        if employee_ids is not None:
            if not employee_ids:
                return []
            placeholders = ",".join("?" for _ in employee_ids)
            query += f" AND employee_id IN ({placeholders})"
            params.extend(employee_ids)
        if practice and str(practice).upper() != "ALL":
            query += " AND LOWER(TRIM(practice)) = LOWER(?)"
            params.append(str(practice).strip())
        if department and str(department).upper() != "ALL":
            query += " AND LOWER(TRIM(department)) = LOWER(?)"
            params.append(str(department).strip())
        query += " ORDER BY employee_name COLLATE NOCASE"
        return [dict(row) for row in connection.execute(query, tuple(params)).fetchall()]
    finally:
        connection.close()


def _safe_filename(value):
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", str(value)).strip("_.") or "attendance"


def _split_name(name):
    parts = str(name or "").strip().split()
    if not parts:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[-1], " ".join(parts[:-1])


def _day_code(day):
    if not day:
        return "A"
    status = str(day.get("status") or "").strip().upper()
    if status == "FULL_DAY":
        return "P"
    if status == "HALF_DAY":
        return "HD"
    if status == "LEAVE" or day.get("leave_type"):
        lt = str(day.get("leave_type") or "").strip().upper()
        if lt in ("", "LEAVE", "L"):
            return "L"
        return lt if len(lt) <= 3 else "L"
    if status == "WEEKEND":
        return "W"
    if status == "HOLIDAY":
        return "H"
    if status == "ABSENT":
        return "A"
    worked = int(day.get("worked_minutes") or 0)
    if worked > 450:
        return "P"
    if worked >= 240:
        return "HD"
    return "A"


# Attendance styles matching template palette
STYLE_FULL_DAY = {
    "fill": PatternFill("solid", fgColor="C6EFCE"),  # soft green
    "font": Font(name="Calibri", size=11, bold=True, color="006100"),
}
STYLE_HALF_DAY = {
    "fill": PatternFill("solid", fgColor="FFF3CD"),  # soft warm amber/yellow
    "font": Font(name="Calibri", size=11, bold=True, color="856404"),
}
STYLE_ABSENT = {
    "fill": PatternFill("solid", fgColor="FFC7CE"),  # soft light red
    "font": Font(name="Calibri", size=11, bold=True, color="9C0006"),
}
STYLE_LEAVE = {
    "fill": PatternFill("solid", fgColor="D9E1F2"),  # soft blue/purple
    "font": Font(name="Calibri", size=11, bold=True, color="1F4E78"),
}
STYLE_WEEK_OFF = {
    "fill": PatternFill("solid", fgColor="FCE4D6"),  # soft warm peach
    "font": Font(name="Calibri", size=11, bold=True, color="854D0E"),
}
STYLE_REMAINING = {
    "fill": PatternFill("solid", fgColor="F2F2F2"),  # neutral light grey
    "font": Font(name="Calibri", size=11, color="595959"),
}
STYLE_UPCOMING = {
    "fill": PatternFill("solid", fgColor="E6E8EB"),  # subtle template grey for future days
    "font": Font(name="Calibri", size=11, bold=True, color="9CA3AF"),
}


def _get_attendance_style(code_or_status):
    norm = str(code_or_status or "").strip().upper()
    if norm in ("P", "FULL_DAY", "FULL DAY", "PRESENT"):
        return STYLE_FULL_DAY
    if norm in ("HD", "HALF_DAY", "HALF DAY"):
        return STYLE_HALF_DAY
    if norm in ("A", "ABSENT"):
        return STYLE_ABSENT
    if norm in ("L", "LEAVE") or "LEAVE" in norm or "LOP" in norm:
        return STYLE_LEAVE
    if norm in ("W", "WO", "WEEKEND", "WEEK OFF", "WEEK_OFF"):
        return STYLE_WEEK_OFF
    return STYLE_REMAINING


def _style_cell(cell, *, fill=None, font=None, horizontal="center"):
    cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=True)
    if fill:
        cell.fill = fill
    if font:
        cell.font = font


def _add_attendance_rows(workbook, employees, start_date, end_date, is_template=False):
    report_month = start_date.replace(day=1)
    last_day = calendar.monthrange(report_month.year, report_month.month)[1]
    today = date.today()

    sheet = workbook.worksheets[0]
    sheet.title = f"FINAL Attendance {report_month.strftime('%b%Y')}"

    # If loaded from template with sample data, clean existing employee data rows (row 3 onwards)
    if is_template and sheet.max_row > 2:
        sheet.delete_rows(3, sheet.max_row - 2)

    dark_fill = PatternFill("solid", fgColor="1F4E78")
    light_fill = PatternFill("solid", fgColor="D9EAF7")
    white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold = Font(name="Calibri", size=11, bold=True)
    thin_gray = Side(style="thin", color="B7C9D6")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    # --------------------------------------------------------------------------
    # Row 1: Report Title / Month in I1, Weekday Headers in J1..AN1, Summary in AO1
    # --------------------------------------------------------------------------
    for col in range(1, 48):
        cell = sheet.cell(1, col)
        cell.fill = dark_fill
        cell.font = white_bold
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.cell(1, 9, datetime(report_month.year, report_month.month, 1))
    sheet.cell(1, 9).number_format = "mmm-yy"
    sheet.cell(1, 41, "Attendance - Summary")

    for day_num in range(1, 32):
        col = 9 + day_num
        if day_num <= last_day:
            day_date = date(report_month.year, report_month.month, day_num)
            sheet.cell(1, col, day_date.strftime("%a"))
        else:
            sheet.cell(1, col).value = None

    sheet.row_dimensions[1].height = 24

    # --------------------------------------------------------------------------
    # Row 2: Fixed Metadata in A2..I2, Dates in J2..AN2, Summary Headers in AO2..AU2
    # --------------------------------------------------------------------------
    for col in range(1, 48):
        cell = sheet.cell(2, col)
        cell.fill = light_fill
        cell.font = bold
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col, label in enumerate(FIXED_COLUMNS, start=1):
        sheet.cell(2, col, label)

    for day_num in range(1, 32):
        col = 9 + day_num
        if day_num <= last_day:
            day_date = date(report_month.year, report_month.month, day_num)
            sheet.cell(2, col, day_date)
            sheet.cell(2, col).number_format = "d"
        else:
            sheet.cell(2, col).value = None
            sheet.cell(2, col).fill = STYLE_REMAINING["fill"]

    for col, label in enumerate(SUMMARY_COLUMNS, start=41):
        sheet.cell(2, col, label)

    sheet.row_dimensions[2].height = 34

    # Pre-calculate month statistics (weekends & working days)
    total_days_in_month = last_day
    weekends_in_month = sum(
        1 for d in range(1, last_day + 1)
        if date(report_month.year, report_month.month, d).weekday() >= 5
    )
    actual_working_days_in_month = total_days_in_month - weekends_in_month

    # --------------------------------------------------------------------------
    # Rows 3+: Dynamic Employee Records
    # --------------------------------------------------------------------------
    for index, employee in enumerate(employees, start=1):
        row_num = index + 2
        last_name, first_name = _split_name(employee.get("employee_name"))
        employee_id = str(employee.get("employee_id") or "")
        department = employee.get("department") or ""
        practice = employee.get("practice") or ""
        reporting_person = employee.get("reports_to") or ""
        work_location = employee.get("work_location") or "Miracle City (HQ)"
        hr_contact = employee.get("hr_contact") or employee.get("hr_email") or ""

        fixed_values = [
            index, employee_id, last_name, first_name, department,
            practice, reporting_person, work_location, hr_contact,
        ]
        for col, value in enumerate(fixed_values, start=1):
            cell = sheet.cell(row_num, col, value)
            cell.border = border
            if col in (1, 2):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        days = _history_for_range(
            employee_id, report_month, date(report_month.year, report_month.month, last_day)
        )
        present_count = half_day_count = leave_count = holiday_count = 0

        for day_num in range(1, 32):
            col = 9 + day_num
            cell = sheet.cell(row_num, col)
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

            if day_num > last_day:
                cell.value = None
                cell.fill = STYLE_REMAINING["fill"]
                cell.font = STYLE_REMAINING["font"]
                continue

            current_date = date(report_month.year, report_month.month, day_num)
            day = days[day_num - 1] if day_num - 1 < len(days) else None

            # Dynamic check: Upcoming / future days within the current month
            is_upcoming = (
                report_month.year == today.year
                and report_month.month == today.month
                and current_date > today
            )
            is_in_requested_range = (start_date <= current_date <= end_date)

            if is_upcoming:
                # Matches template: future days display 0 with soft neutral grey fill
                cell.value = 0
                cell.fill = STYLE_UPCOMING["fill"]
                cell.font = STYLE_UPCOMING["font"]
            elif not is_in_requested_range:
                cell.value = ""
                cell.fill = STYLE_REMAINING["fill"]
                cell.font = STYLE_REMAINING["font"]
            else:
                code = _day_code(day)
                cell.value = code
                style = _get_attendance_style(code)
                cell.fill = style["fill"]
                cell.font = style["font"]

                if code == "P":
                    present_count += 1
                elif code == "HD":
                    half_day_count += 1
                elif (day or {}).get("status") == "LEAVE" or code in ("L", "Leave") or bool((day or {}).get("leave_type")):
                    leave_count += 1
                elif (day or {}).get("status") == "HOLIDAY" or code == "H":
                    holiday_count += 1

        # Summary columns (AO..AU / Cols 41..47)
        summary_values = [
            present_count,
            half_day_count,
            actual_working_days_in_month - holiday_count,
            leave_count,
            holiday_count,
            weekends_in_month,
            total_days_in_month,
        ]
        for s_idx, value in enumerate(summary_values, start=41):
            cell = sheet.cell(row_num, s_idx, value)
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Column Widths: metadata columns A-I, uniform date columns J-AN (all 5.0), summary AO-AU
    widths = {
        1: 5.0, 2: 14.0, 3: 16.0, 4: 18.0, 5: 20.0, 6: 20.0, 7: 22.0, 8: 22.0, 9: 18.0,
    }
    for col in range(10, 41):
        widths[col] = 5.0
    widths[41] = 5.0   # P
    widths[42] = 8.0   # Half Day
    widths[43] = 12.0  # Actual Working Days
    widths[44] = 18.0  # Tot. Leaves
    widths[45] = 12.0  # Holidays
    widths[46] = 12.0  # Week ends
    widths[47] = 12.0  # No of Days in a Month

    for col, width in widths.items():
        sheet.column_dimensions[get_column_letter(col)].width = width

    sheet.freeze_panes = "J3"
    sheet.auto_filter.ref = f"A2:AU{max(2, len(employees) + 2)}"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "$1:$2"

    # --------------------------------------------------------------------------
    # Secondary Sheet: "Attendance Summary"
    # --------------------------------------------------------------------------
    if "Attendance Summary" in workbook.sheetnames:
        summary = workbook["Attendance Summary"]
        if summary.max_row > 1:
            summary.delete_rows(2, summary.max_row - 1)
    else:
        summary = workbook.create_sheet("Attendance Summary")
        summary.append([
            "Date", "Employee ID", "Employee Name", "Shift Start", "Shift End",
            "First In", "Last Out", "Worked Hours", "Worked Minutes", "Break Duration",
            "Status", "Leave Type", "Leave Reason",
        ])
        for cell in summary[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F2937")
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for employee in employees:
        emp_id = str(employee.get("employee_id") or "")
        emp_days = _history_for_range(emp_id, start_date, end_date)
        for offset, day in enumerate(emp_days):
            if day is None:
                continue
            cur_date = start_date + timedelta(days=offset)
            is_flex = bool(employee.get("is_flexible_shift"))
            shift_start_val = "Flexible" if is_flex else (employee.get("shift_start") or "")
            shift_end_val = "Flexible" if is_flex else (employee.get("shift_end") or "")
            status_val = day.get("status") or "ABSENT"
            summary.append([
                cur_date,
                emp_id,
                employee.get("employee_name") or "",
                shift_start_val,
                shift_end_val,
                day.get("first_in") or "",
                day.get("last_out") or "",
                day.get("worked_hours") or "0h 0m",
                int(day.get("worked_minutes") or 0),
                day.get("break_duration") or "0h 0m",
                status_val,
                day.get("leave_type") or "",
                day.get("reason") or "",
            ])
            row_idx = summary.max_row
            status_cell = summary.cell(row_idx, 11)
            status_style = _get_attendance_style(status_val)
            status_cell.fill = status_style["fill"]
            status_cell.font = status_style["font"]
            status_cell.alignment = Alignment(horizontal="center", vertical="center")

    workbook.active = sheet


def _create_report(employees, period, selected_date, report_name, start_date=None, end_date=None):
    start_date_obj, end_date_obj = _period_dates(period, selected_date, start_date, end_date)
    workbook, is_template = _load_or_create_workbook()
    _add_attendance_rows(workbook, employees, start_date_obj, end_date_obj, is_template=is_template)
    output = BytesIO()
    workbook.save(output)
    filename = (
        f"{_safe_filename(report_name)}_{period}_attendance_"
        f"{start_date_obj:%Y%m%d}_to_{end_date_obj:%Y%m%d}.xlsx"
    )
    return output.getvalue(), filename


def create_employee_report(employee_id, period="month", selected_date=None, start_date=None, end_date=None, date=None):
    employee = get_employee(employee_id)
    if employee is None:
        raise LookupError(f"Employee '{employee_id}' not found")
    target_date = selected_date or date
    return _create_report([dict(employee)], period, target_date, employee_id, start_date, end_date)


def create_all_employees_report(period="month", selected_date=None, start_date=None, end_date=None, department=None, date=None, employee_ids=None, practice=None):
    target_date = selected_date or date
    return _create_report(
        _employee_rows(department=department, employee_ids=employee_ids, practice=practice), period, target_date, "all_employees", start_date, end_date
    )
